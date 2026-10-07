"""QA local bằng identity giả; không chứa raw user thật."""
import csv
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import run_reports as reports


def fixture(path, bad_fix=False, missing_mode=False, nested=False, nullable_json=False, boundaries=False):
    import duckdb
    con = duckdb.connect()
    con.execute('''CREATE TABLE fixture(app_id VARCHAR,user_pseudo_id VARCHAR,event_ts TIMESTAMP,
        event_name VARCHAR,event_params VARCHAR,user_properties VARCHAR,mode_fixed VARCHAR,
        first_clear_650_ts TIMESTAMP,app_version VARCHAR,country VARCHAR,platform VARCHAR)''')
    rows = []
    def add(user, name, day, level=3, ab='0', params=None, property_level=None, mode='classic', cutoff=None):
        ep = dict(level=level, mode='classic', phase=1)
        ep.update(params or {})
        up = dict(level=level if property_level is None else property_level, mode='classic', phase=1, firebase_exp_abt_22=ab)
        rows.append(('id6758755718', 'fake_' + user, day + ' 01:00:00', name, json.dumps(ep), json.dumps(up),
                     mode, cutoff, '9.9.9', 'Vietnam', 'ios'))
    add('a', 'level_start', '2026-10-02')
    add('b', 'level_start', '2026-10-02', property_level=4)
    add('c', 'level_start', '2026-10-02')
    add('d', 'level_start', '2026-10-04')
    for completion in (5,70):
        add('a', 'level_end_turn', '2026-10-02', params=dict(start_count=1,success=False,completion=completion))
    add('b', 'level_end_turn', '2026-10-02', params=dict(start_count=2,success=False,completion=5))
    add('c', 'level_end_turn', '2026-10-02', params=dict(start_count=1,success=False,completion=20))
    add('b', 'level_end', '2026-10-02', params=dict(start_count=2,success=False,completion=50,continue_times=0))
    add('c', 'session_start', '2026-10-06')
    add('a', 'level_start_turn', '2026-10-02', params=dict(start_count=1,coin=100,ticket=50,booster_x=2,use_booster_count=17))
    add('b', 'level_start_turn', '2026-10-02', params=dict(start_count=2,coin=10))
    add('a', 'spend_virtual_currency', '2026-10-02', params=dict(virtual_currency_name='coin',value_game_currency=10))
    add('a', 'spend_virtual_currency', '2026-10-02', params=dict(virtual_currency_name='booster_x',virtual_currency_type='booster',value=1,value_game_currency=35))
    add('a', 'paid_ad_impression', '2026-10-02', params=dict(ad_format='interstitial'))
    add('a', 'paid_ad_impression', '2026-10-02', params=dict(ad_format='banner'))
    add('wide', 'level_start', '2026-10-02', level=300, ab='2')
    add('phase', 'level_end_turn', '2026-10-02', level=4,
        params=dict(phase=2,start_count=1,success=False,completion=25))
    cutoff = '2026-04-01 01:00:00'
    add('old', 'level_end', '2026-04-01', level=650, ab='1', params=dict(success=True), cutoff=cutoff)
    add('old', 'level_start', '2026-10-02', ab='1', mode='classic' if bad_fix else 'sl', cutoff=cutoff)
    add('old', 'paid_ad_impression', '2026-10-02', ab='1', mode=None, cutoff=cutoff, params=dict(ad_format='interstitial',mode=None))
    if boundaries:
        def edge(user, ts, name='level_start', level=400):
            add(user,name,ts[:10],level=level)
            rows[-1]=rows[-1][:2]+(ts,)+rows[-1][3:]
        edge('before','2026-10-01 16:59:59')
        edge('at_start','2026-10-01 17:00:00')
        edge('last_report','2026-10-05 16:59:59',level=401)
        edge('after_report','2026-10-05 17:00:00',level=402)
        edge('at_start','2026-10-06 16:59:59',name='session_start')
        edge('outside_observation','2026-10-06 17:00:00',name='session_start')
        edge('future','2026-10-12 01:00:00',name='session_start')
    con.executemany('INSERT INTO fixture VALUES (?,?,?,?,?,?,?,?,?,?,?)', rows)
    if nullable_json:
        # An actual activity event must survive missing params/properties.
        con.execute("UPDATE fixture SET event_params=NULL, user_properties=NULL WHERE event_name='session_start'")
        # Irrelevant app_loading rows: historical/in-window NULL, empty and malformed JSON.
        con.execute('''INSERT INTO fixture
            SELECT app_id,user_pseudo_id,ts,'app_loading',ep,up,mode_fixed,
                   first_clear_650_ts,app_version,country,platform
            FROM (SELECT * FROM fixture WHERE event_name='level_start' LIMIT 1) f
            CROSS JOIN (VALUES
                (TIMESTAMP '2026-05-01',NULL::VARCHAR,NULL::VARCHAR),
                (TIMESTAMP '2026-10-02','', '{bad json'),
                (TIMESTAMP '2026-10-02',NULL::VARCHAR,'{}')) x(ts,ep,up)''')
    columns = '* EXCLUDE(mode_fixed)' if missing_mode else '*'
    if nested:
        schemas=con.execute('SELECT json_group_structure(event_params::JSON), json_group_structure(user_properties::JSON) FROM fixture').fetchone()
        columns='* EXCLUDE(event_params,user_properties),from_json(event_params,'+reports.literal(schemas[0])+') AS event_params,from_json(user_properties,'+reports.literal(schemas[1])+') AS user_properties'
    con.execute('COPY (SELECT ' + columns + ' FROM fixture) TO ' + reports.literal(path) + ' (FORMAT PARQUET)')
    con.close()


class ReportsTest(unittest.TestCase):
    def execute(self, tmp, sql_mode="optimized", config=None, work_dir=None, **fixture_args):
        src = Path(tmp) / 'raw_fixed.parquet'
        fixture(src, **fixture_args)
        dest = Path(tmp) / 'out'
        args = SimpleNamespace(input=str(src),output=str(dest),config=str(config or ROOT/'config.json'),sql_mode=sql_mode,
                               memory_limit='256MB',threads=1,row_cap=1000,work_dir=work_dir)
        reports.run(args)
        return dest

    def test_outputs_formulas_censoring_and_modes(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            dest = self.execute(tmp)
            def read(name):
                with (dest/name).open(encoding='utf-8-sig') as f:
                    return list(csv.DictReader(f))
            dm = read('data_metrics.csv')
            self.assertEqual(len(dm),2)  # corrected SL excluded; all levels includes300
            self.assertEqual({(r['ab_group'],float(r['level'])) for r in dm},{('0',3),('2',300)})
            d = dm[0]
            self.assertEqual(d['ab_group'],'0')
            self.assertEqual(int(d['user_start_count']),4)
            self.assertEqual(float(d['completion_rate_first_attempt_pct']),50)
            self.assertEqual(float(d['coin_balance_per_user']),50)
            self.assertEqual(float(d['total_coin_balance_per_user']),76)
            self.assertEqual(float(d['coin_spend_per_user']),2.5)
            self.assertEqual(float(d['total_coin_spend_per_user']),11.25)
            self.assertAlmostEqual(float(d['imp_per_lau']),1/3)  # excludes banner
            self.assertAlmostEqual(float(d['inter_imp_per_lau']),1/3)
            self.assertEqual(d['rwd_per_lau'],'N/A')
            self.assertEqual(d['churn_rate_d7_pct'],'N/A')
            first=read('loss_first_attempt.csv')[0]
            alls=read('loss_all_attempts.csv')[0]
            self.assertEqual((int(first['user_count']),int(first['lose_count'])),(2,3))
            self.assertEqual((int(alls['user_count']),int(alls['lose_count'])),(3,4))
            self.assertEqual(first['churn_users_d3'], '1')
            self.assertEqual(first['churn_users_d3'],alls['churn_users_d3'])
            self.assertAlmostEqual(float(first['churn_rate_d3_pct']),100/3)
            self.assertEqual(float(first['0-10']),50)
            self.assertEqual(float(read('loss_first_attempt.csv')[1]['level']),4.2)  # Loss keepsphase from sourceBI
            wb=load_workbook(dest/'sandjam_reports.xlsx')
            self.assertEqual(wb.sheetnames,['Data Metrics','Loss 1st attempt','Loss All'])
            self.assertTrue(wb['Loss All'].conditional_formatting)
            self.assertEqual(wb['Data Metrics'].freeze_panes,'C2')
            manifest=json.loads((dest/'manifest.json').read_text())
            for path,expected in manifest['files'].items():
                self.assertEqual(reports.sha(dest/path),expected)
            self.assertEqual(manifest['status'],'complete')
            self.assertEqual(len(list(dest.glob('*.csv'))),3)
            self.assertNotIn('fake_', (dest/'data_readme.md').read_text())

    def test_null_and_malformed_json_keep_rows_and_activity(self):
        import duckdb
        with tempfile.TemporaryDirectory() as tmp1,tempfile.TemporaryDirectory() as tmp2:
            baseline=self.execute(tmp1)
            tolerant=self.execute(tmp2,nullable_json=True)
            for filename,_,_ in reports.OUTPUTS.values():
                self.assertEqual((baseline/filename).read_bytes(),(tolerant/filename).read_bytes())
            qa=json.loads((tolerant/'manifest.json').read_text())['source_qa']
            self.assertEqual(qa['json_policy'],'tolerant_null_no_global_gate')
            self.assertNotIn('invalid_json',qa)
            con=duckdb.connect()
            cfg=json.loads((ROOT/'config.json').read_text())
            args=SimpleNamespace(input=str(Path(tmp2)/'raw_fixed.parquet'))
            reports.prepare_input(con,args,cfg)
            self.assertEqual(con.execute('SELECT COUNT(*) FROM raw_fixed_input').fetchone()[0],qa['rows'])
            self.assertEqual(con.execute("SELECT COUNT(*) FROM raw_fixed_input WHERE event_name='app_loading'").fetchone()[0],2)
            self.assertEqual(con.execute("SELECT COUNT(*) FROM raw_fixed_input WHERE event_name='session_start' AND event_params IS NULL AND user_properties IS NULL").fetchone()[0],1)
            con.close()

    def test_optimized_legacy_parity_timezone_and_d7(self):
        for nested,obs,asof,denom,versions in (
                (False,'2026-10-06','2026-10-07','property',[]),
                (True,'2026-10-06','2026-10-07','event',[]),
                (False,'2026-10-12','2026-10-13','property',['9.9.9'])):
            with self.subTest(nested=nested,obs=obs,denom=denom), tempfile.TemporaryDirectory() as tmp1, tempfile.TemporaryDirectory() as tmp2:
                cfg=json.loads((ROOT/'config.json').read_text())
                cfg.update(observation_end=obs,as_of_date=asof,ads_denominator=denom,versions=versions)
                config=Path(tmp1)/'config.json';config.write_text(json.dumps(cfg))
                fixture_args=dict(boundaries=True,nested=nested,nullable_json=not nested)
                fast=self.execute(tmp1,config=config,**fixture_args)
                old=self.execute(tmp2,sql_mode='legacy',config=config,**fixture_args)
                for filename,_,_ in reports.OUTPUTS.values():
                    self.assertEqual((fast/filename).read_bytes(),(old/filename).read_bytes())
                with (fast/'data_metrics.csv').open(encoding='utf-8-sig') as f:
                    data=list(csv.DictReader(f))
                self.assertTrue(any(float(r['level'])==400 for r in data))
                self.assertTrue(any(float(r['level'])==401 for r in data))
                self.assertFalse(any(float(r['level'])==402 for r in data))
                self.assertEqual(data[0]['churn_rate_d7_pct']=='N/A',obs=='2026-10-06')
                meta=json.loads((fast/'manifest.json').read_text())
                self.assertEqual(meta['sql_mode'],'optimized')
                self.assertEqual(meta['source_qa']['parquet_data_scan_statements'],1)
                self.assertTrue(meta['sql_timings'])
                self.assertFalse((fast/'_work').exists())
                self.assertFalse((fast/'_spill').exists())
                plan=json.loads((fast/'sql_executed/source_scan_plan.json').read_text())
                def nodes(items):
                    for item in items:
                        yield item
                        yield from nodes(item.get('children',[]))
                scans=[n for n in nodes(plan) if n['name'].strip()=='READ_PARQUET']
                self.assertEqual(len(scans),1)
                self.assertIn('event_ts',str(scans[0]['extra_info']))

    def test_optimized_source_detached_after_normalization(self):
        import duckdb
        from lib.sql_statements import split_sql
        with tempfile.TemporaryDirectory() as tmp:
            src=Path(tmp)/'input.parquet';fixture(src,nullable_json=True,boundaries=True)
            con=duckdb.connect();con.execute("SET TimeZone='UTC'")
            cfg=json.loads((ROOT/'config.json').read_text())
            con.execute('CREATE MACRO date_sub_days(d,n) AS CAST(d AS DATE)-CAST(n AS INTEGER)')
            con.execute('CREATE MACRO date_add_days(d,n) AS CAST(d AS DATE)+CAST(n AS INTEGER)')
            qa=reports.prepare_input(con,SimpleNamespace(input=str(src)),cfg)
            self.assertEqual(qa['utc_start'],'2026-10-01 17:00:00')
            self.assertEqual(qa['utc_end_exclusive'],'2026-10-06 17:00:00')
            self.assertFalse(con.execute("SELECT COUNT(*) FROM raw_fixed_input WHERE event_ts<TIMESTAMP '2026-10-01 17:00:00' OR event_ts>=TIMESTAMP '2026-10-06 17:00:00'").fetchone()[0])
            con.execute(reports.config_view(cfg))
            for filename,text in reports.render_sql(cfg):
                for statement in split_sql(text):
                    con.execute(statement)
                if filename=='01_normalize.sql':
                    src.unlink()  # Remaining stages succeed even without the Parquet file.
            for view in ('levelplay_report','loss_report'):
                plan=con.execute('EXPLAIN SELECT * FROM '+view).fetchone()[1]
                self.assertNotIn('READ_PARQUET',plan)
                self.assertTrue(con.execute('SELECT COUNT(*) FROM '+view).fetchone()[0])
            con.close()

    def test_string_timestamp_input_matches_native(self):
        import duckdb
        with tempfile.TemporaryDirectory() as tmp:
            tmp=Path(tmp);native=self.execute(tmp)
            con=duckdb.connect()
            con.execute("COPY (SELECT * REPLACE(event_ts::VARCHAR AS event_ts,first_clear_650_ts::VARCHAR AS first_clear_650_ts) FROM read_parquet("+reports.literal(tmp/'raw_fixed.parquet')+")) TO "+reports.literal(tmp/'strings.parquet')+" (FORMAT PARQUET)")
            con.close()
            for mode in ('optimized','legacy'):
                dest=tmp/('strings_'+mode)
                args=SimpleNamespace(input=str(tmp/'strings.parquet'),output=str(dest),config=str(ROOT/'config.json'),
                    memory_limit='256MB',threads=1,row_cap=1000,sql_mode=mode)
                reports.run(args)
                for filename,_,_ in reports.OUTPUTS.values():
                    self.assertEqual((native/filename).read_bytes(),(dest/filename).read_bytes())
                if mode=='optimized':
                    meta=json.loads((dest/'manifest.json').read_text())
                    self.assertFalse(meta['source_qa']['native_timestamp_filter'])

    def test_external_work_directory_cleanup(self):
        with tempfile.TemporaryDirectory() as tmp:
            work=Path(tmp)/'ssd'
            self.execute(tmp,work_dir=str(work))
            self.assertFalse(list(work.iterdir()))
        with tempfile.TemporaryDirectory() as tmp:
            work=Path(tmp)/'ssd'
            with self.assertRaises(ValueError):
                self.execute(tmp,work_dir=str(work),bad_fix=True)
            self.assertFalse(list(work.iterdir()))
            self.assertFalse((Path(tmp)/'out').exists())

    def test_input_must_be_corrected(self):
        for kwargs in (dict(missing_mode=True),dict(bad_fix=True)):
            with self.subTest(kwargs=kwargs),tempfile.TemporaryDirectory() as tmp:
                with self.assertRaises(ValueError):
                    self.execute(tmp,**kwargs)
                self.assertFalse((Path(tmp)/'out').exists())

    def test_existing_output_not_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            dest=self.execute(tmp)
            marker=dest/'manifest.json'
            digest=reports.sha(marker)
            with self.assertRaises(ValueError):
                self.execute(tmp)
            self.assertEqual(reports.sha(marker),digest)

    def test_struct_json_same_outputs(self):
        with tempfile.TemporaryDirectory() as tmp1,tempfile.TemporaryDirectory() as tmp2:
            string_output=self.execute(tmp1)
            struct_output=self.execute(tmp2,nested=True)
            for filename,_,_ in reports.OUTPUTS.values():
                self.assertEqual((string_output/filename).read_bytes(),(struct_output/filename).read_bytes())


if __name__=='__main__':
    unittest.main()
