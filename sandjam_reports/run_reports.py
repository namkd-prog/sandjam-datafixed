#!/usr/bin/env python3
"""Chạy ba report trên Parquet raw_fixed của Nam, không gọi StarRocks."""
import argparse
import csv
import hashlib
import json
import re
import shutil
import tempfile
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

from lib.sql_statements import split_sql

ROOT = Path(__file__).resolve().parent
OUTPUTS = {
    'levelplay_report': ('data_metrics.csv', 'Data Metrics', None),
    'loss_first': ('loss_first_attempt.csv', 'Loss 1st attempt', 'first'),
    'loss_all': ('loss_all_attempts.csv', 'Loss All', 'all'),
}
REQUIRED = {
    'app_id', 'user_pseudo_id', 'event_ts', 'event_name', 'event_params',
    'user_properties', 'mode_fixed', 'first_clear_650_ts', 'app_version', 'country', 'platform',
}


def literal(value):
    return "'" + str(value).replace("'", "''") + "'"


def array(values):
    return '[' + ','.join(literal(v) for v in values) + ']::VARCHAR[]'


def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def validate_config(cfg):
    for key in ('report_start', 'report_end', 'as_of_date', 'observation_end'):
        date.fromisoformat(cfg[key])
    if cfg['report_start'] > cfg['report_end'] or cfg['report_end'] > cfg['observation_end']:
        raise ValueError('Cần report_start <= report_end <= observation_end.')
    if cfg['observation_end'] >= cfg['as_of_date']:
        raise ValueError('observation_end phải là ngày hoàn chỉnh trước as_of_date.')
    if cfg['ab_groups'] != ['0', '1', '2']:
        raise ValueError('Gói này chia AB22 nhóm 0/1/2; sửa có chủ đích khi tái dùng task khác.')
    if cfg['ads_denominator'] not in ('property', 'event'):
        raise ValueError('ads_denominator: property hoặc event.')
    for key in ('versions', 'coin_balance_resources'):
        if not isinstance(cfg[key], list) or any(not isinstance(v, str) for v in cfg[key]):
            raise ValueError(key + ' phải là list string.')


def render_sql(cfg):
    values = {'app_id': literal(cfg['app_id']),
              'variants': ','.join(literal(v) for v in cfg['ab_groups']),
              'ads_denominator': 'ps.users' if cfg['ads_denominator'] == 'property' else 'c.user_start_count'}
    result = []
    for path in sorted((ROOT / 'sql').glob('*.sql')):
        text = path.read_text(encoding='utf-8')
        for key, value in values.items():
            text = text.replace('{{' + key + '}}', value)
        if '{{' in text:
            raise ValueError('SQL còn placeholder: ' + path.name)
        # Cache chỉ các cửa sổ report/observation; không giữ toàn history trong RAM.
        for view in ('ab_classic_events', 'metric_events', 'daily_sessions', 'churn_play_pre'):
            text = text.replace('CREATE OR REPLACE TEMP VIEW ' + view + ' AS',
                                'CREATE OR REPLACE TEMP TABLE ' + view + ' AS')
        result.append((path.name, text))
    return result


def config_view(cfg):
    return f'''CREATE OR REPLACE TEMP VIEW report_config AS SELECT
        DATE {literal(cfg['report_start'])} AS report_start,
        DATE {literal(cfg['report_end'])} AS report_end,
        DATE {literal(cfg['as_of_date'])} AS as_of_date,
        DATE {literal(cfg['observation_end'])} AS observation_end,
        {array(cfg['versions'])} AS versions,
        {array(cfg['coin_balance_resources'])} AS coin_balance_resources'''


def prepare_input(con, args, cfg):
    path = Path(args.input).expanduser()
    if path.is_dir():
        pattern = str(path / '**/*.parquet')
    else:
        pattern = str(path)
    con.execute('CREATE TEMP VIEW parquet_source AS SELECT * FROM read_parquet(' + literal(pattern) + ', union_by_name=true, hive_partitioning=true)')
    types = {row[0]: row[1] for row in con.execute('DESCRIBE parquet_source').fetchall()}
    missing = REQUIRED - types.keys()
    if missing:
        raise ValueError('Input phải là raw_fixed Parquet đầy đủ; thiếu: ' + ', '.join(sorted(missing)))
    if any(types[k] in ('BIGINT', 'INTEGER', 'DOUBLE', 'UBIGINT') for k in ('event_ts', 'first_clear_650_ts')):
        raise ValueError('Timestamp phải là UTC timestamp/ISO string, không dùng epoch integer.')
    selected = []
    for name in sorted(REQUIRED):
        quoted = '"' + name + '"'
        if name in ('event_params', 'user_properties'):
            typ = types[name]
            if typ.startswith(('STRUCT', 'MAP')):
                expr = 'to_json(' + quoted + ')::VARCHAR'
            elif typ in ('VARCHAR', 'JSON'):
                # Missing/malformed JSON is read as NULL; preserve the input row.
                expr = 'TRY_CAST(' + quoted + ' AS JSON)::VARCHAR'
            else:
                raise ValueError(name + ' cần JSON string/Map/Struct; GA4 arrays cần adapter.')
        elif name in ('event_ts', 'first_clear_650_ts'):
            expr = '(TRY_CAST(' + quoted + " AS TIMESTAMPTZ) AT TIME ZONE 'UTC')"
        else:
            expr = quoted + '::VARCHAR'
        selected.append(expr + ' AS ' + quoted)
    con.execute('CREATE TEMP VIEW raw_fixed_input AS SELECT ' + ','.join(selected) + ' FROM parquet_source')
    check = con.execute(f'''SELECT COUNT(*) AS rows,
      COUNT(*) FILTER (WHERE event_ts IS NULL) AS invalid_timestamp,
      MIN(DATE(event_ts+INTERVAL 7 HOURS)),MAX(DATE(event_ts+INTERVAL 7 HOURS))
      FROM raw_fixed_input WHERE app_id={literal(cfg['app_id'])}''').fetchone()
    if not check[0] or check[1]:
        raise ValueError(f'Input rỗng hoặc timestamp không hợp lệ: rows={check[0]}, invalid_timestamp={check[1]}.')
    if str(check[2]) > cfg['report_start'] or str(check[3]) < cfg['observation_end']:
        raise ValueError('Input không phủ report/observation_end. Cần toàn bộ raw_fixed, gồm session tới observation_end.')
    bad_fix = con.execute(f'''SELECT COUNT(*) FROM raw_fixed_input
      WHERE app_id={literal(cfg['app_id'])} AND COALESCE(mode_fixed,'')<>'sl'
      AND json_extract_string(event_params,'$.mode')='classic'
      AND event_ts>first_clear_650_ts''').fetchone()[0]
    if bad_fix:
        raise ValueError(f'Có {bad_fix} event original classic sau cutoff nhưng mode_fixed chưa là sl; kiểm lại bước process của Nam.')
    return dict(rows=check[0], min_date=str(check[2]), max_date=str(check[3]),
                invalid_timestamp=check[1], json_policy='tolerant_null_no_global_gate',
                inconsistent_mode_fix=bad_fix,
                input_schema={k: types[k] for k in sorted(REQUIRED)})


def export_csv(con, view, kind, destination, row_cap):
    if kind:
        cursor = con.execute('SELECT * EXCLUDE(loss_kind) FROM loss_report WHERE loss_kind=? ORDER BY ab_group,level', [kind])
    else:
        cursor = con.execute('SELECT * FROM levelplay_report ORDER BY ab_group,level')
    headers = [x[0] for x in cursor.description]
    if any('user_pseudo_id' == x for x in headers):
        raise ValueError('Report không được xuất identity.')
    rows = 0
    with destination.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        while batch := cursor.fetchmany(500):
            rows += len(batch)
            if rows > row_cap:
                raise ValueError('Report vượt row_cap; tăng --row-cap, không cắt kết quả.')
            writer.writerows([['N/A' if v is None else v for v in row] for row in batch])
    return {'rows': rows, 'columns': headers, 'sha256': sha(destination)}


def excel_report(folder):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for filename, sheet_name, _ in OUTPUTS.values():
        ws = wb.create_sheet(sheet_name)
        with (folder / filename).open(encoding='utf-8-sig', newline='') as f:
            for n, row in enumerate(csv.reader(f)):
                if n == 0:
                    ws.append(row)
                    continue
                values = []
                for i, value in enumerate(row):
                    if i == 0 or value == 'N/A':
                        values.append(value)
                    else:
                        try:
                            number = Decimal(value)
                            values.append(int(number) if number == number.to_integral_value() else float(number))
                        except Exception:
                            values.append(value)
                ws.append(values)
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 48
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='355D88')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for i, cell in enumerate(ws[1], 1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = 13 if i <= 2 else 22
            for column_cell in list(ws.columns)[i - 1][1:]:
                is_count = str(cell.value) in ('user_start_count','user_count','lose_count','churn_users_d3')
                column_cell.number_format = '#,##0' if is_count else ('#,##0.00' if i > 2 else '0.##')
            if i <= 2 or ws.max_row < 2:
                continue
            is_bucket = re.fullmatch(r'\d+-\d+', str(cell.value))
            red = str(cell.value).startswith('churn')
            rule = ColorScaleRule(start_type='num' if is_bucket else 'min', start_value=0 if is_bucket else None,
                                  start_color='EFF6FF', end_type='num' if is_bucket else 'max',
                                  end_value=100 if is_bucket else None, end_color='E88B8B' if red else '347FD0')
            ws.conditional_formatting.add(f'{letter}2:{letter}{ws.max_row}', rule)
    wb.save(folder / 'sandjam_reports.xlsx')


def data_readme(cfg, source_qa, outputs, windows):
    return f'''# Dữ liệu Sandjam iOS

App `{cfg['app_id']}`. Report **{cfg['report_start']}–{cfg['report_end']}**, ngày UTC+7, bao gồm hai biên.
Input là **Parquet raw_fixed đã được process của Nam sửa**, không đọc/query StarRocks. Gameplay lọc `mode_fixed='classic'`;
ads/IAP lấy `user_properties.mode` và áp cùng `first_clear_650_ts` để loại classic sau cutoff. Không sửa JSON hay level ID.
JSON NULL/rỗng/sai cú pháp không chặn job: khi đọc, giá trị không parse được trở thành NULL; giữ nguyên dòng và Parquet gốc. Không quét toàn history để kiểm JSON. Metric cần field nào thì áp điều kiện field đó; session vẫn được giữ theo logic observation.
Chia theo `firebase_exp_abt_22`, nhóm 0/1/2. Version: {', '.join(cfg['versions']) or 'mọi version'}. **Mọi level**, không giới hạn1–200.

| File / tab Excel | Nội dung |
|---|---|
| data_metrics.csv / Data Metrics | Data Metrics của Level Play, AB × level |
| loss_first_attempt.csv / Loss 1st attempt | level_end_turn, success=false, start_count=1, completion nonnull |
| loss_all_attempts.csv / Loss All | Cùng điều kiện loss, không giới hạn start_count |

`sandjam_reports.xlsx` có đúng 3 tab, freeze AB/level, filter và heatmap. CSV UTF-8 BOM.
Giá trị phần trăm là **percentage points**:16.67 nghĩa16.67%, không phải0.1667. CSV giữ precision tính toán; Excel hiển thị2 số lẻ.
`N/A` là NULL/thiếu observation hoặc không có format ads;0 là số tính được bằng0.

## Công thức và phạm vi kiểm

- User count=distinct user; lose count=số event. Vì có lặp event/start_count, Loss1st cũng có thể lose_count>user_count.
- Distribution: distinct user trong từng bucket / distinct user thua tại level. All attempts có thể cộng bucket >100% vì user xuất hiện ở nhiều bucket. Biên `<10`, `<20`,...; bucket cuối>=90.
- Ordinary churn=startersX không startX+1 trong cùng period/AB. D3/D7 Data Metrics=Non Return, highest level theo user/ngày, không session_start/screen_view trong+1..N; denominator là starters trong cohort mature.
- Completion Rate First Attempt=AVG completion trên level_end, continue_times 0/null; không phải win-first rate và không thêm start_count=1.
- Coin spend=spend coin value_game_currency / starters; Total coin spend=mọi spend value_game_currency / starters.
- Balance numerator: resources không âm ở level_start_turn attempt1; coin hoặc mọi coin/ticket/booster hợp lệ trừ use_booster_count. Denominator=distinct start_turn user mọi attempt.
- Total IMP/LAU chỉ Inter+Rewarded. Cả ba cột ads dùng **{cfg['ads_denominator']}-level starters**; property là candidate khớp cả 3 cột ở 10 level ảnh, chưa có SQL BI xác nhận. Có thể chọn event trong config để so mẫu số event-level trong code Nam.
- Loss Churn: starters mature D3, highest level user/ngày; không level_end cùng level trong report, không session_start/screen_view sau ngày chơi tới observation_end. Churn độc lập user thua/attempts, nên cùng số ở2 bảng Loss và có thể lớn hơn user_count.
- **Loss Churn là numeric reconstruction** đã khớp count/rate ở24 level 3–26 trên raw BI 23/09–06/10. SQL phụ BI/Long chưa xác nhận business logic. Không gọi nó là công thức nguồn BI đã xác nhận.
- QA ảnh: 240 ô nhìn thấy (2 attempts × 24 level × user_count/lose_count/0–10/Churn count/rate) khớp; 144 ô Loss 1st đủ 10 bucket level 3–14 đã khớp trước đó. 130 ô Data Metrics ngoài ads và 30 ô ads dùng mẫu số property ở level 1–10 khớp. Coverage này không chứng nhận mọi level/report 02–05/10 đã sửa mode.
- Không áp Internal/pre-publish theo yêu cầu người dùng. Có system filters từ SQL nguồn, classic strict, phase<=1 cho Data Metrics; Loss distribution giữphase theo SQL show-query.

## Observation và censoring

As-of: {cfg['as_of_date']}; ngày hoàn chỉnh cuối nguồn: {cfg['observation_end']}. Ngày eligible cuối:
{json.dumps(windows, ensure_ascii=False)}.
Với default 07/10 và observation 06/10: D3 dùng play_date đến03/10; D7 đến29/09. Report02–05/10 chỉ có02–03/10 mature D3;
D7 không có cohort đủ tuổi nên **N/A**, không đổi thành 0. Muốn D7 của toànperiod cần input có activity tới 12/10 và as-of 13/10 trở đi.
MAX ngày nguồn chỉ là bound, không chứng minh partition của observation_end đã đầy đủ; DE phải xác nhận ngày hoàn chỉnh.
Không suy thiếu event là thiếu hành vi; không dedupe event tự ý để khớp BI.

Input QA aggregate:{json.dumps({k:v for k,v in source_qa.items() if k!='input_schema'},ensure_ascii=False)}.
Số dòng:{json.dumps({k:v['rows'] for k,v in outputs.items()},ensure_ascii=False)}.
Chi tiết cấu hình, schema, nguồn SQL và hash CSV/Excel ởmanifest.json; không xuất identifier.
'''


def run(args):
    import duckdb
    cfg = json.loads(Path(args.config).read_text(encoding='utf-8'))
    validate_config(cfg)
    dest = Path(args.output).expanduser().resolve()
    if dest.exists():
        raise ValueError('Output đã tồn tại; chọn thư mục mới, không ghi đè snapshot.')
    dest.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix='.' + dest.name + '-', dir=dest.parent))
    con = None
    try:
        con = duckdb.connect()
        con.execute("SET TimeZone='UTC'")
        con.execute('SET memory_limit=?', [args.memory_limit])
        con.execute('SET threads=?', [args.threads])
        con.execute('SET temp_directory=?', [str(staging / '_spill')])
        con.execute('CREATE MACRO date_sub_days(d,n) AS CAST(d AS DATE)-CAST(n AS INTEGER)')
        con.execute('CREATE MACRO date_add_days(d,n) AS CAST(d AS DATE)+CAST(n AS INTEGER)')
        source_qa = prepare_input(con, args, cfg)
        con.execute(config_view(cfg))
        rendered = render_sql(cfg)
        sql_dir = staging / 'sql_executed'
        sql_dir.mkdir()
        (sql_dir / 'config.sql').write_text(config_view(cfg) + ';\n')
        for filename, text in rendered:
            (sql_dir / filename).write_text(text, encoding='utf-8')
            print('SQL ' + filename, flush=True)
            for number, statement in enumerate(split_sql(text), 1):
                try:
                    con.execute(statement)
                except Exception as exc:
                    raise RuntimeError(f'{filename}, statement{number}: {exc}') from exc
        windows = [{'n': n, 'eligible_end': str(end)} for n,end in con.execute('SELECT n,eligible_end FROM churn_windows ORDER BY n').fetchall()]
        outputs = {}
        for view, (filename, sheet, kind) in OUTPUTS.items():
            outputs[filename] = export_csv(con, view, kind, staging / filename, args.row_cap)
            print(f'{filename}: {outputs[filename]["rows"]} dòng', flush=True)
        if not outputs['data_metrics.csv']['rows']:
            raise ValueError('Data Metrics rỗng: kiểm mode_fixed, period và AB; không coi run này là hoàn tất.')
        con.close()
        shutil.rmtree(staging / '_spill', ignore_errors=True)
        excel_report(staging)
        (staging / 'data_readme.md').write_text(data_readme(cfg, source_qa, outputs, windows), encoding='utf-8')
        manifest = dict(status='complete', created_at_utc=datetime.now(timezone.utc).isoformat(),
                        engine='duckdb', duckdb_version=duckdb.__version__, input=str(args.input),
                        config=cfg, source_qa=source_qa, churn_windows=windows, outputs=outputs,
                        source_hashes={str(p.relative_to(ROOT)):sha(p) for p in ROOT.rglob('*')
                                       if p.is_file() and '__pycache__' not in p.parts and p.suffix in ('.py','.sql','.json','.txt')},
                        files={str(p.relative_to(staging)):sha(p) for p in staging.rglob('*') if p.is_file()})
        (staging / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        staging.rename(dest)
        print('COMPLETE ' + str(dest), flush=True)
    except Exception:
        if con is not None:
            con.close()
        shutil.rmtree(staging, ignore_errors=True)
        raise


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--input', required=True, help='Thư mục raw_fixed hoặc glob Parquet; giữ đầy đủ history/session, không chỉ gameplay classic.')
    p.add_argument('--output', required=True, help='Thư mục local mới cho3 CSV/Excel/readme/manifest.')
    p.add_argument('--config', default=str(ROOT / 'config.json'))
    p.add_argument('--memory-limit', default='8GB', help='DuckDB spill xuống temp khi vượt memory_limit.')
    p.add_argument('--threads', type=int, default=4)
    p.add_argument('--row-cap', type=int, default=20000)
    args = p.parse_args()
    if args.threads < 1 or args.row_cap < 1:
        p.error('threads và row-cap phải dương.')
    run(args)


if __name__ == '__main__':
    main()
